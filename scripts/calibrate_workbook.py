"""Build the Grid 5 workbook: Grid 4 with its coefficients recalibrated.

Grid 4 overstated the footprint. On the JSL Jajpur route it gave ~4.6 tCO2e/t
with Scope 1+2 at ~3.2, against the ~1.8-2.2 tCO2e/tcs Scope 1+2 intensity
Jindal Stainless publishes, and ~34 GJ/t where a melt-to-coil stainless route
runs at roughly 10-15 GJ/t. Almost all of the excess traces to electricity:
every row's kWh/t was backed out of an illustrative Scope 2 figure, and the
route summed to ~3,300 kWh/t where published stainless-route benchmarks put it
nearer 1,000-1,200 kWh/t.

Nothing structural changes here. The sheets, headers, rows, departments,
processes, variations, the x/y, a-g and p/q variables, the formula shapes and
the grid emission factors are all Grid 4's. Only the per-row coefficients move,
and every change is recorded in the row's own notes and in a "Calibration Basis"
sheet so the old and new values can be compared side by side.

How a row is recalibrated
-------------------------
Each row has three primary coefficients per charge (virgin / scrap):

* ``kwh`` - electricity per tonne, which feeds Scope 2;
* ``s1``  - direct (Scope 1) emissions per tonne;
* ``s3``  - upstream (Scope 3, Cat. 1-8) emissions per tonne.

A calibration entry either sets a new value outright (``Set``) or scales the
Grid 4 value (``Scale``). The dependent columns follow from those:

* CO2, CH4 and N2O are the direct-gas breakdown, so they scale with Scope 1;
* HFC, PFC, SF6 and NF3 are refrigerant and switchgear losses, which scale with
  the electrical plant, so they follow kWh;
* specific energy (final energy, GJ/t) is rebuilt as
  ``0.0036 * kWh + fuel``, with fuel = Scope 1 / fuel emission factor
  (diesel 0.074 tCO2/GJ for yard and despatch equipment, natural gas
  0.056 tCO2/GJ elsewhere). Rows whose Scope 1 is process carbon rather than
  fuel (EAF, converters, BF) carry an explicit specific energy instead.

The two transport rows are left exactly as Grid 4 has them: their per-tonne
haulage figures are already in the range of published rail and road freight
intensities.

Usage::

    python scripts/calibrate_workbook.py      # writes Grid_5.xlsx
    python scripts/extract_workbook.py        # regenerates data/carbon_grid.json
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "source" / "Stainless_Steel_Carbon_Accounting_Grid_4.xlsx"
OUTPUT = ROOT / "data" / "source" / "Stainless_Steel_Carbon_Accounting_Grid_5.xlsx"

GRID_SHEET = "Carbon Accounting Grid"
HEADER_ROWS = 4
FIRST_METRIC_COLUMN = 15  # 1-based column O
NOTES_COLUMN = 26  # 1-based column Z
METRIC_KEYS = ("scope1", "scope2", "scope3", "co2", "ch4", "n2o", "hfc", "pfc", "sf6", "nf3", "sec")

#: The reference grid factor Grid 4 used to back its kWh/t out of Scope 2.
REFERENCE_GRID_FACTOR = 0.77
KWH_TO_GJ = 0.0036
DIESEL_EF = 0.074  # tCO2 per GJ
GAS_EF = 0.056  # tCO2 per GJ


# --------------------------------------------------------------------------- #
# Calibration table
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class Set:
    """New (virgin, scrap) values, replacing Grid 4's."""

    virgin: float
    scrap: float


@dataclass(frozen=True)
class Scale:
    """Multiply Grid 4's (virgin, scrap) values."""

    virgin: float
    scrap: Optional[float] = None

    def pair(self) -> Tuple[float, float]:
        return self.virgin, self.virgin if self.scrap is None else self.scrap


Spec = object  # Set | Scale | None


@dataclass(frozen=True)
class Row:
    reason: str
    kwh: Spec = None
    s1: Spec = None
    s3: Spec = None
    sec: Optional[Set] = None


# Scope 3 of the metal charge. Upstream burden of the virgin units dominates a
# stainless footprint (ferrochrome, nickel units, manganese, pig iron). Grid 4
# put it at 1.1 tCO2e per tonne of virgin charge, below what those units carry
# even for low-nickel grades; scrap enters burden-free under the cut-off
# approach and carries only collection and processing. The values below are for
# JSL's grade mix (200, 300 and 400 series), which sits well under a 304-only
# charge. Every melting variation carries the same charge burden, because the
# upstream footprint belongs to what is charged, not to the furnace melting it.
CHARGE_S3 = Set(3.20, 0.20)

CALIBRATION: Dict[int, Row] = {
    # ---------------- RMHS ----------------
    2: Row("Inspection is already near-negligible; unchanged."),
    3: Row("Inspection is already near-negligible; unchanged."),
    4: Row(
        "Scrap shredding and baling draws ~20-25 kWh per tonne of scrap; yard diesel "
        "for mobile handling ~3-5 L/t.",
        kwh=Scale(0.6), s1=Scale(0.35),
    ),
    5: Row(
        "Sizing and screening: ~4-8 kWh/t and light mobile-plant diesel.",
        kwh=Scale(0.35), s1=Scale(0.4),
    ),
    6: Row(
        "Stockyard stacking and reclaiming: ~3-5 kWh/t; mobile plant diesel.",
        kwh=Scale(0.3), s1=Scale(0.4),
    ),
    7: Row("Charge-bucket make-up: cranes and weighing, ~2-4 kWh/t.", kwh=Scale(0.3)),
    8: Row(
        "Belt conveying is ~0.5-1 kWh per tonne per km; in-plant runs total ~3-6 kWh/t.",
        kwh=Scale(0.25),
    ),
    # ---------------- Melt Shop ----------------
    9: Row("Scrap charging cranes and magnets: ~4-10 kWh/t.", kwh=Scale(0.35)),
    10: Row(
        "Stainless EAF melts scrap at ~420-480 kWh/t; a virgin-heavy charge (ferroalloys, "
        "pig iron, DRI) needs more, ~600-700 kWh/t. Direct emissions from electrodes "
        "(~2 kg/t), charge carbon and burners are ~0.06-0.15 tCO2/t.",
        kwh=Set(650, 460), s1=Set(0.14, 0.07), s3=CHARGE_S3, sec=Set(3.1, 2.1),
    ),
    11: Row(
        "Induction furnaces melt scrap at ~550-650 kWh/t with negligible direct emissions.",
        kwh=Set(750, 600), s3=CHARGE_S3, sec=Set(2.8, 2.2),
    ),
    12: Row(
        "Integrated BF/converter: Scope 1 already carries the iron-making reductant "
        "(~1.8 tCO2/t hot metal), so Scope 3 carries only the alloy units; electricity "
        "for blast, oxygen and auxiliaries ~250-350 kWh/t.",
        kwh=Set(300, 65), s3=Set(1.80, 0.20), sec=Set(19.0, 3.5),
    ),
    13: Row(
        "Vacuum induction melting runs ~800-1,000 kWh/t.",
        kwh=Set(1000, 800), s3=CHARGE_S3, sec=Set(3.8, 3.1),
    ),
    14: Row(
        "Vacuum arc remelting runs ~1,000-1,300 kWh/t on top of primary melting.",
        kwh=Set(1300, 1100), s3=CHARGE_S3, sec=Set(4.9, 4.2),
    ),
    15: Row(
        "Electroslag remelting runs ~1,200-1,600 kWh/t.",
        kwh=Set(1500, 1300), s3=CHARGE_S3, sec=Set(5.6, 4.9),
    ),
    16: Row(
        "AOD needs little electricity of its own; oxygen, argon and nitrogen supply plus "
        "fume extraction come to ~40-60 kWh/t. Decarburisation CO2 left as in Grid 4.",
        kwh=Set(60, 45), sec=Set(0.9, 0.6),
    ),
    17: Row(
        "VOD vacuum pumping and gas supply: ~60-100 kWh/t.",
        kwh=Set(100, 70), sec=Set(1.2, 0.9),
    ),
    18: Row("K-OBM-S: converter blowing, gas supply and fume system ~50-70 kWh/t.",
            kwh=Set(70, 50), sec=Set(1.3, 0.9)),
    19: Row("CLU: as AOD with steam dilution, ~50-65 kWh/t.",
            kwh=Set(65, 48), sec=Set(1.0, 0.7)),
    20: Row(
        "Ladle furnace arcing is ~30-50 kWh/t; trim-alloy additions ~0.04-0.08 tCO2e/t "
        "upstream.",
        kwh=Set(55, 35), s1=Set(0.006, 0.004), s3=Set(0.08, 0.04),
    ),
    21: Row("Argon stirring: gas supply energy only.", kwh=Scale(0.5)),
    22: Row("Nitrogen stirring: gas supply energy only.", kwh=Scale(0.5)),
    23: Row("VD/RH degassing with mechanical pumps: ~30-45 kWh/t.", kwh=Set(45, 30)),
    24: Row(
        "Slab casters use ~20-35 kWh/t including the water system; tundish preheating "
        "and torch cutting are small.",
        kwh=Set(35, 28), s1=Scale(0.4), s3=Scale(0.7),
    ),
    25: Row("Vertical caster: as the curved-mould caster, slightly more.",
            kwh=Set(40, 32), s1=Scale(0.4), s3=Scale(0.7)),
    26: Row("Thin-slab casting: ~35-45 kWh/t with the direct-link furnace excluded.",
            kwh=Set(45, 36), s1=Scale(0.4), s3=Scale(0.7)),
    27: Row("SEN: a refractory practice, not a load of its own.", kwh=Scale(0.5)),
    28: Row("Argon shrouding: gas supply energy only.", kwh=Scale(0.5)),
    29: Row("Mould EMS draws ~5-10 kWh/t.", kwh=Set(8, 6)),
    30: Row("Soft reduction: segment actuation, ~3-5 kWh/t.", kwh=Set(4, 3)),
    31: Row("Air-mist cooling: spray pumps, ~5-10 kWh/t.", kwh=Set(8, 6)),
    # ---------------- Hot Rolling ----------------
    32: Row("Slab grinding: ~6-10 kWh/t.", kwh=Scale(0.25)),
    33: Row(
        "Walking-beam reheating burns ~1.2-1.5 GJ/t of gas for stainless; fans ~5-10 kWh/t.",
        kwh=Scale(0.4), s1=Set(0.075, 0.07),
    ),
    34: Row("HP descaling pumps: ~6-10 kWh/t.", kwh=Scale(0.3)),
    35: Row("Roughing mill: ~30-40 kWh/t.", kwh=Scale(0.3)),
    36: Row("Crop shear and coil box / edge heating: ~5-10 kWh/t and little fuel.",
            kwh=Scale(0.3), s1=Scale(0.35)),
    37: Row("Finishing train: ~45-60 kWh/t for stainless strip.", kwh=Scale(0.28)),
    38: Row("Laminar cooling pumps: ~5-8 kWh/t.", kwh=Scale(0.2)),
    39: Row("Down-coiler or plate shear: ~5-8 kWh/t.", kwh=Scale(0.18)),
    # ---------------- Annealing ----------------
    40: Row("Gas-fired solution annealing: ~0.8-1.1 GJ/t fuel, ~20-30 kWh/t.",
            kwh=Scale(0.65), s1=Scale(0.9)),
    41: Row("Subcritical anneal: lower soak temperature.", kwh=Scale(0.65), s1=Scale(0.9)),
    42: Row("Full anneal / temper: batch furnace.", kwh=Scale(0.65), s1=Scale(0.9)),
    43: Row("Bright annealing in hydrogen / cracked ammonia.", kwh=Scale(0.8), s1=Scale(0.8)),
    44: Row("Open annealing with pickling.", kwh=Scale(0.65), s1=Scale(0.9)),
    45: Row("Continuous strand / tube annealing.", kwh=Scale(0.65), s1=Scale(0.9)),
    46: Row("Induction heating to solution temperature takes ~250-320 kWh/t.",
            kwh=Set(300, 280)),
    # ---------------- Descaling ----------------
    47: Row("HP spray descaling: ~8-12 kWh/t.", kwh=Scale(0.35)),
    48: Row("Shot blasting: ~20-30 kWh/t.", kwh=Scale(0.45)),
    49: Row("Scale breaker: ~8-12 kWh/t.", kwh=Scale(0.35)),
    50: Row("Salt bath: gas-heated, electricity ~15-20 kWh/t.", kwh=Scale(0.45), s1=Scale(0.8)),
    51: Row("Ultrasonic agitation: ~10-15 kWh/t.", kwh=Scale(0.35)),
    # ---------------- Pickling ----------------
    52: Row("Tank pickling pumps, heating and acid regeneration auxiliaries: ~20-30 kWh/t.",
            kwh=Scale(0.65)),
    53: Row("Batch pickling.", kwh=Scale(0.65)),
    54: Row("Paste pickling: manual, minimal energy; unchanged."),
    55: Row("Spray pickling.", kwh=Scale(0.65)),
    # ---------------- Cold Rolling ----------------
    56: Row("20-Hi Sendzimir: ~120-160 kWh/t depending on total reduction.", kwh=Scale(0.7)),
    57: Row("18-Hi cluster mill.", kwh=Scale(0.7)),
    58: Row("6-Hi UC mill.", kwh=Scale(0.7)),
    59: Row("Tandem cold mill.", kwh=Scale(0.7)),
    60: Row("CAPL: gas-fired anneal ~0.9-1.0 GJ/t plus pickling; electricity kept.",
            s1=Scale(0.9)),
    61: Row("Skin-pass mill: ~10-15 kWh/t.", kwh=Scale(0.45)),
    62: Row("Bright annealing in the cold-rolling sequence.", kwh=Scale(0.8), s1=Scale(0.8)),
    63: Row("Tension levelling and slitting: ~10-15 kWh/t.", kwh=Scale(0.4)),
    # ---------------- Outbound ----------------
    64: Row("Packing: strapping and wrapping machines, ~5-8 kWh/t.", kwh=Scale(0.6)),
    65: Row("Warehousing: lighting and cranes, ~4-6 kWh/t.", kwh=Scale(0.25)),
    66: Row("Picking: forklifts and cranes, ~2-4 kWh/t.", kwh=Scale(0.4)),
}


# --------------------------------------------------------------------------- #
# Formula reading and writing
# --------------------------------------------------------------------------- #
_NUM = r"[0-9]+(?:\.[0-9]+)?(?:[eE][-+]?[0-9]+)?"
_LINEAR = re.compile(rf"x\*({_NUM})\+y\*({_NUM})")


def fmt(value: float) -> str:
    """Plain decimal, six significant figures, no exponent — as Grid 4 writes them."""
    if value == 0:
        return "0"
    text = f"{float(f'{value:.6g}'):.12f}".rstrip("0").rstrip(".")
    return text or "0"


def read_pair(formula: str) -> Tuple[float, float]:
    match = _LINEAR.search(formula)
    if match is None:
        raise ValueError(f"not a linear x/y formula: {formula}")
    return float(match.group(1)), float(match.group(2))


def write_pair(formula: str, virgin: float, scrap: float) -> str:
    return _LINEAR.sub(f"x*{fmt(virgin)}+y*{fmt(scrap)}", formula, count=1)


def resolve(spec: Spec, old: Tuple[float, float]) -> Tuple[float, float]:
    if spec is None:
        return old
    if isinstance(spec, Set):
        return spec.virgin, spec.scrap
    mv, ms = spec.pair()
    return old[0] * mv, old[1] * ms


def ratio(new: float, old: float) -> float:
    return new / old if old else 1.0


# --------------------------------------------------------------------------- #
# Notes
# --------------------------------------------------------------------------- #
_NOTE_NAMES = {
    "EF_S1": "scope1", "EF_S3": "scope3", "CO2": "co2", "CH4": "ch4", "N2O": "n2o",
    "HFC": "hfc", "PFC": "pfc", "SF6": "sf6", "NF3": "nf3", "SEC": "sec",
}


def rewrite_note(note: str, department: str, values: Dict[str, Tuple[float, float]],
                 kwh: Tuple[float, float], old: Dict[str, Tuple[float, float]],
                 old_kwh: Tuple[float, float], reason: str) -> str:
    def sub_value(text: str, name: str, suffix: str, value: float) -> str:
        pattern = re.compile(rf"\b{name}_{suffix}=({_NUM})")
        if not pattern.search(text):
            raise ValueError(f"note has no {name}_{suffix}")
        return pattern.sub(f"{name}_{suffix}={fmt(value)}", text, count=1)

    text = note
    for name, key in _NOTE_NAMES.items():
        virgin, scrap = values[key]
        long_form = f"{name}_virgin=" in text
        text = sub_value(text, name, "virgin" if long_form else "v", virgin)
        text = sub_value(text, name, "scrap" if long_form else "s", scrap)
    # EF_S2 in the note is the Scope 2 at the 0.77 kg/kWh reference.
    text = sub_value(text, "EF_S2", "virgin", kwh[0] * REFERENCE_GRID_FACTOR / 1000)
    text = sub_value(text, "EF_S2", "scrap", kwh[1] * REFERENCE_GRID_FACTOR / 1000)
    text, count = re.subn(
        rf"= ?{_NUM} kWh/t \(virgin\), {_NUM} kWh/t \(scrap\)",
        f"= {fmt(kwh[0])} kWh/t (virgin), {fmt(kwh[1])} kWh/t (scrap)",
        text,
    )
    if count != 1:
        raise ValueError("note has no kWh line")
    text = re.sub(
        r"^Coefficients used in this row's formulas \(illustrative,",
        "Coefficients used in this row's formulas (Grid 5 calibration,",
        text,
    )
    changed = []
    for label, new_pair, old_pair in (
        ("kWh/t", kwh, old_kwh),
        ("Scope 1", values["scope1"], old["scope1"]),
        ("Scope 3", values["scope3"], old["scope3"]),
        ("SEC GJ/t", values["sec"], old["sec"]),
    ):
        if any(abs(n - o) > 1e-12 for n, o in zip(new_pair, old_pair)):
            changed.append(
                f"{label} {fmt(old_pair[0])}/{fmt(old_pair[1])} → "
                f"{fmt(new_pair[0])}/{fmt(new_pair[1])}"
            )
    summary = "; ".join(changed) if changed else "no change"
    text += (
        f"\nGrid 5 calibration: {reason} Grid 4 → Grid 5 (virgin/scrap): {summary}."
    )
    return text


# --------------------------------------------------------------------------- #
# Build
# --------------------------------------------------------------------------- #
def calibrate_row(row_cells, calibration: Row, department: str):
    formulas = {
        key: str(row_cells[FIRST_METRIC_COLUMN - 1 + offset].value).strip()
        for offset, key in enumerate(METRIC_KEYS)
    }
    old = {key: read_pair(formula) for key, formula in formulas.items()}
    old_kwh = old["scope2"]  # the scope2 formula's coefficients are kWh/t

    kwh = resolve(calibration.kwh, old_kwh)
    s1 = resolve(calibration.s1, old["scope1"])
    s3 = resolve(calibration.s3, old["scope3"])

    s1_ratio = (ratio(s1[0], old["scope1"][0]), ratio(s1[1], old["scope1"][1]))
    kwh_ratio = (ratio(kwh[0], old_kwh[0]), ratio(kwh[1], old_kwh[1]))

    new = {"scope1": s1, "scope2": kwh, "scope3": s3}
    for key in ("co2", "ch4", "n2o"):
        new[key] = (old[key][0] * s1_ratio[0], old[key][1] * s1_ratio[1])
    for key in ("hfc", "pfc", "sf6", "nf3"):
        new[key] = (old[key][0] * kwh_ratio[0], old[key][1] * kwh_ratio[1])

    untouched = calibration.kwh is None and calibration.s1 is None and calibration.sec is None
    if calibration.sec is not None:
        new["sec"] = (calibration.sec.virgin, calibration.sec.scrap)
    elif untouched:
        new["sec"] = old["sec"]
    else:
        fuel_ef = DIESEL_EF if department in ("RMHS", "Outbound") else GAS_EF
        new["sec"] = tuple(KWH_TO_GJ * k + s / fuel_ef for k, s in zip(kwh, s1))

    for offset, key in enumerate(METRIC_KEYS):
        cell = row_cells[FIRST_METRIC_COLUMN - 1 + offset]
        cell.value = write_pair(formulas[key], *new[key])

    note_values = {key: new[key] for key in _NOTE_NAMES.values()}
    note_cell = row_cells[NOTES_COLUMN - 1]
    note_cell.value = rewrite_note(
        note_cell.value or "", department, note_values, kwh, old, old_kwh, calibration.reason
    )
    return old, new


def build(source: Path = SOURCE, output: Path = OUTPUT) -> list:
    workbook = openpyxl.load_workbook(source)
    sheet = workbook[GRID_SHEET]
    log = []
    process_id = 0
    for row_cells in sheet.iter_rows(min_row=HEADER_ROWS + 1):
        if row_cells[0].value is None or row_cells[FIRST_METRIC_COLUMN - 1].value is None:
            continue
        process_id += 1
        transport = str(row_cells[12].value).strip() == "p"
        if transport:
            log.append((process_id, row_cells, None, None, "Transport row: unchanged."))
            continue
        calibration = CALIBRATION.get(process_id)
        if calibration is None:
            raise SystemExit(f"row {process_id} has no calibration entry")
        old, new = calibrate_row(row_cells, calibration, row_cells[0].value)
        log.append((process_id, row_cells, old, new, calibration.reason))

    # Title and description: say which edition this is, keep everything else.
    sheet["A1"].value = (
        "Carbon Accounting Formula Database — Indian Stainless Steel Industry (Grid 5, calibrated)"
    )
    sheet["A2"].value = (
        str(sheet["A2"].value)
        + " GRID 5: coefficients recalibrated against published stainless-route benchmarks "
        "(see the 'Calibration Basis' sheet); rows, headers, variables and formula shapes are "
        "unchanged from Grid 4."
    )
    _write_basis_sheet(workbook, log)
    workbook.save(output)
    return log


BENCHMARKS = (
    ("Jindal Stainless, Scope 1+2 GHG intensity, FY23", "2.15 tCO2e/tcs",
     "Company disclosure; includes captive power and captive ferrochrome, which this grid "
     "has no rows for."),
    ("Jindal Stainless, Scope 1+2 GHG intensity, FY26", "1.76 tCO2e/tcs",
     "Company disclosure (18% below FY23)."),
    ("Jindal Stainless, scrap in the charge, FY26", "70.12%",
     "Company disclosure, company-wide."),
    ("worldstainless / ISSF, Scope 1 average", "~0.4 tCO2/t stainless",
     "Scrap-based producers; 80% of producers in 0.2-0.6."),
    ("worldstainless / ISSF, Scope 2 average", "~0.4-0.5 tCO2/t stainless",
     "Global average, on grids cleaner than India's."),
    ("worldstainless / ISSF, Scope 3 vs scrap share", "2.90 / 1.65 / 1.15 tCO2/t at 50 / 75 / 85% scrap",
     "Austenitic (Ni-bearing) basis; linear in scrap share. Lower-nickel 200/400 series sit below it."),
    ("Stainless EAF electricity", "~420-480 kWh/t (scrap), ~600-700 kWh/t (virgin-heavy)",
     "Typical published range for stainless EAF melting."),
    ("Melt-to-cold-coil electricity", "~1,000-1,200 kWh/t",
     "EAF, AOD, LF, caster, hot strip mill, Z-mill, CAPL and auxiliaries."),
)


def _write_basis_sheet(workbook, log) -> None:
    name = "Calibration Basis"
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    wrap = Alignment(wrap_text=True, vertical="top")

    sheet["A1"].value = "Calibration Basis — Grid 4 → Grid 5"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"].value = (
        "Grid 5 keeps Grid 4's structure exactly and recalibrates per-row coefficients. "
        "Grid 4's electricity intensities summed to ~3,300 kWh/t on the JSL routes, about three "
        "times published stainless-route figures, which put Scope 1+2 near 3.2 tCO2e/t against "
        "JSL's disclosed 1.76-2.15. Scope 3 of the virgin charge moves the other way (1.1 → "
        "3.2 tCO2e/t virgin charge), because ferroalloy and nickel units carry more upstream "
        "carbon than Grid 4 assumed. Trace gases follow Scope 1 (CO2, CH4, N2O) or electricity "
        "(HFC, PFC, SF6, NF3); SEC is final energy = 0.0036 x kWh + fuel. Still not plant-measured."
    )
    sheet["A2"].alignment = wrap
    sheet.merge_cells("A2:L2")
    sheet.row_dimensions[2].height = 90

    sheet["A4"].value = "Benchmark"
    sheet["B4"].value = "Value"
    sheet["C4"].value = "Note"
    for cell in (sheet["A4"], sheet["B4"], sheet["C4"]):
        cell.font = bold
        cell.fill = head_fill
    row = 5
    for label, value, note in BENCHMARKS:
        sheet.cell(row, 1, label).alignment = wrap
        sheet.cell(row, 2, value).alignment = wrap
        sheet.cell(row, 3, note).alignment = wrap
        row += 1

    row += 1
    headers = (
        "#", "Department", "Technique / Process", "Variation",
        "Grid 4 kWh/t (v)", "Grid 5 kWh/t (v)", "Grid 4 kWh/t (s)", "Grid 5 kWh/t (s)",
        "Grid 4 Scope 1 (v/s)", "Grid 5 Scope 1 (v/s)", "Grid 4 Scope 3 (v/s)", "Grid 5 Scope 3 (v/s)",
        "Grid 4 SEC (v/s)", "Grid 5 SEC (v/s)", "Reason",
    )
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row, col, header)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = wrap
    row += 1
    for process_id, cells, old, new, reason in log:
        base = [process_id, cells[0].value, cells[1].value, cells[2].value]
        if old is None:
            values = ["", "", "", "", "", "", "", "", "", ""]
        else:
            def pair(p):
                return f"{fmt(p[0])} / {fmt(p[1])}"
            values = [
                float(fmt(old["scope2"][0])), float(fmt(new["scope2"][0])),
                float(fmt(old["scope2"][1])), float(fmt(new["scope2"][1])),
                pair(old["scope1"]), pair(new["scope1"]),
                pair(old["scope3"]), pair(new["scope3"]),
                pair(old["sec"]), pair(new["sec"]),
            ]
        for col, value in enumerate(base + values + [reason], start=1):
            sheet.cell(row, col, value).alignment = wrap
        row += 1

    widths = (5, 13, 30, 26, 11, 11, 11, 11, 15, 15, 15, 15, 14, 14, 70)
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 60


def main() -> int:
    log = build()
    changed = sum(1 for entry in log if entry[2] is not None)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}: {changed} rows recalibrated, "
          f"{len(log) - changed} transport rows unchanged.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

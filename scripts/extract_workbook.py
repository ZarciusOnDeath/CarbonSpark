"""Regenerate ``data/carbon_grid.json`` from the source workbook.

The application never reads the .xlsx at runtime — it reads the JSON this script
produces, so the parsing assumptions live in one reviewable place. Re-run it
whenever the workbook is updated::

    python scripts/extract_workbook.py [path/to/workbook.xlsx]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "data" / "source" / "Stainless_Steel_Carbon_Accounting_Grid_3.xlsx"
OUTPUT = ROOT / "data" / "carbon_grid.json"

GRID_SHEET = "Carbon Accounting Grid"
MIX_SHEET = "Grid Energy Mix Factors"
DOWNSTREAM_SHEET = "Downstream Scope 3 (Product)"

#: Columns M-W of the grid sheet, in order.
METRIC_KEYS = (
    "scope1",
    "scope2",
    "scope3",
    "co2",
    "ch4",
    "n2o",
    "hfc",
    "pfc",
    "sf6",
    "nf3",
    "sec",
)
FIRST_METRIC_COLUMN = 12  # zero-based index of column M
NOTES_COLUMN = 23  # column X
HEADER_ROWS = 4  # title, description, blank, header


def extract(workbook_path: Path) -> dict:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)

    processes = []
    for row in workbook[GRID_SHEET].iter_rows(min_row=HEADER_ROWS + 1):
        values = [cell.value for cell in row]
        # Trailing disclaimer rows have a department but no formulas.
        if values[0] is None or values[FIRST_METRIC_COLUMN] is None:
            continue
        processes.append(
            {
                "id": len(processes) + 1,
                "department": values[0],
                "process": values[1],
                "variation": values[2],
                "formulas": {
                    key: str(values[FIRST_METRIC_COLUMN + offset]).strip()
                    for offset, key in enumerate(METRIC_KEYS)
                },
                "notes": values[NOTES_COLUMN] or "",
            }
        )

    grid_factors = {}
    for row in workbook[MIX_SHEET].iter_rows(min_row=HEADER_ROWS):
        try:
            factor = float(row[2].value)
        except (TypeError, ValueError):
            continue  # header and note rows
        grid_factors[str(row[1].value).strip()] = {
            "source": row[0].value,
            "ef": factor,
            "basis": row[3].value,
        }

    downstream = [
        {"category": row[0].value, "applicability": row[1].value, "notes": row[2].value}
        for row in workbook[DOWNSTREAM_SHEET].iter_rows(min_row=HEADER_ROWS)
        if row[0].value and row[1].value
    ]

    return {"processes": processes, "grid_factors": grid_factors, "downstream": downstream}


def main() -> int:
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    data = extract(workbook_path)
    if len(data["processes"]) != 71:
        raise SystemExit(
            f"Expected 71 process steps, found {len(data['processes'])} — check the sheet layout."
        )
    if len(data["grid_factors"]) != 7:
        raise SystemExit(
            f"Expected 7 grid sources, found {len(data['grid_factors'])} — check the sheet layout."
        )
    OUTPUT.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
    print(
        f"Wrote {OUTPUT.relative_to(ROOT)}: {len(data['processes'])} processes, "
        f"{len(data['grid_factors'])} grid sources, {len(data['downstream'])} downstream categories."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
